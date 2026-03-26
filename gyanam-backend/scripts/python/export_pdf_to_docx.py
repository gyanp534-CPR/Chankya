from __future__ import annotations

import argparse
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Inches


def export_pdf_to_docx(pdf_path: Path, docx_path: Path, dpi: int, image_width_in: float) -> None:
    doc = fitz.open(pdf_path)
    document = Document()
    document.add_heading("Question Paper (Extracted Pages)", level=1)

    for index, page in enumerate(doc, start=1):
        pix = page.get_pixmap(dpi=dpi)
        img_path = docx_path.parent / f"{docx_path.stem}_page_{index}.png"
        pix.save(img_path)

        document.add_heading(f"Page {index}", level=2)
        document.add_picture(str(img_path), width=Inches(image_width_in))

    document.save(docx_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export PDF pages to a DOCX with images.")
    parser.add_argument("--pdf", required=True, help="Path to the input PDF file.")
    parser.add_argument("--out", required=True, help="Path to the output DOCX file.")
    parser.add_argument("--dpi", type=int, default=300, help="DPI for rendered images.")
    parser.add_argument("--width", type=float, default=6.0, help="Image width in inches.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    pdf_path = Path(args.pdf).expanduser().resolve()
    docx_path = Path(args.out).expanduser().resolve()

    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF not found: {pdf_path}")

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    export_pdf_to_docx(pdf_path, docx_path, args.dpi, args.width)
    print(f"Saved DOCX: {docx_path}")


if __name__ == "__main__":
    main()
