from __future__ import annotations

import argparse
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Inches


def export_pdf_to_docx(pdf_path: Path, docx_path: Path, dpi: int, image_width_in: float) -> None:
    doc = fitz.open(pdf_path)
    document = Document()
    document.add_heading("Question Paper (Extracted Pages)", level=1)

    for index, page in enumerate(doc, start=1):
        pix = page.get_pixmap(dpi=dpi)
        img_path = docx_path.parent / f"{docx_path.stem}_page_{index}.png"
        pix.save(img_path)

        document.add_heading(f"Page {index}", level=2)
        document.add_picture(str(img_path), width=Inches(image_width_in))

    document.save(docx_path)


def find_pdf(year_dir: Path) -> Path | None:
    direct = sorted(year_dir.glob("*.pdf"))
    if direct:
        return direct[0]
    raw_dir = year_dir / "raw"
    raw_pdfs = sorted(raw_dir.glob("*.pdf"))
    return raw_pdfs[0] if raw_pdfs else None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export GS1 PDFs to DOCX files.")
    parser.add_argument("--base", required=True, help="Base papers folder (e.g. data/pyq/papers).")
    parser.add_argument("--out", required=True, help="Output folder for DOCX files.")
    parser.add_argument("--start", type=int, default=2016, help="Start year.")
    parser.add_argument("--end", type=int, default=2025, help="End year.")
    parser.add_argument("--dpi", type=int, default=300, help="DPI for rendered images.")
    parser.add_argument("--width", type=float, default=6.0, help="Image width in inches.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    base_dir = Path(args.base).expanduser().resolve()
    out_dir = Path(args.out).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    for year in range(args.start, args.end + 1):
        year_dir = base_dir / str(year) / "gs1"
        pdf_path = find_pdf(year_dir)

        if not pdf_path:
            print(f"[SKIP] {year}: no GS1 PDF found in {year_dir}")
            continue

        docx_path = out_dir / f"gs1-{year}.docx"
        export_pdf_to_docx(pdf_path, docx_path, args.dpi, args.width)
        print(f"[OK] {year}: {docx_path}")


if __name__ == "__main__":
    main()
